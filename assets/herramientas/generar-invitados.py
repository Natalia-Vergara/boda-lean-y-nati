#!/usr/bin/env python3
"""Arma todo lo que depende de la lista de invitados.

`lista.json` es la única fuente: se toca ahí y se vuelve a correr esto.

    python3 assets/herramientas/generar-invitados.py

Escribe:
  · invitados.js                       — en el repositorio, lo usa el sitio
  · salida/envios.html                 — la página con los mensajes para mandar
  · salida/invitados-lean-y-nati.xlsx  — la planilla

Los códigos son permanentes: una vez que un link salió por WhatsApp, cambiar
el código lo rompe. Si hay que corregir un nombre, se cambia el nombre y el
código queda como está.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
SALIDA = AQUI / 'salida'
SITIO = 'https://natalia-vergara.github.io/boda-lean-y-nati/'

BORDO = '61152A'
TINTA = '3D2E2A'
PAPEL = 'FBF8F2'
LINEA = 'E4D9C6'


def leer_lista():
    datos = json.loads((AQUI / 'lista.json').read_text(encoding='utf-8'))
    codigos = [c for c, _, _ in datos]
    repetidos = {c for c in codigos if codigos.count(c) > 1}
    if repetidos:
        raise SystemExit(f'hay códigos repetidos: {sorted(repetidos)}')
    return [(c, n, int(p)) for c, n, p in datos]


def mensaje(nombre, codigo, pases):
    """El texto listo para pegar en WhatsApp.

    En singular cuando va una sola persona: un mensaje en plural a quien
    viene solo se nota que salió de una plantilla.
    """
    link = f'{SITIO}?i={codigo}'
    if pases == 1:
        cuerpo = (
            f'Nos casamos el viernes 27 de noviembre y nos encantaría que nos acompañes.\n\n'
            f'Acá está nuestra invitación — abrila, que tiene sorpresa:\n{link}\n\n'
            f'Tenés un lugar reservado. Te pedimos confirmar antes del 25 de octubre, '
            f'que es cuando le entregamos la lista definitiva a la finca.\n\n'
            f'¡Te esperamos!\n')
    else:
        cuerpo = (
            f'Nos casamos el viernes 27 de noviembre y nos encantaría que nos acompañen.\n\n'
            f'Acá está nuestra invitación — abrila, que tiene sorpresa:\n{link}\n\n'
            f'Tienen {pases} lugares reservados. Les pedimos confirmar antes del 25 de '
            f'octubre, que es cuando le entregamos la lista definitiva a la finca.\n\n'
            f'¡Los esperamos!\n')
    return f'¡Hola, {nombre}! 🤍\n\n{cuerpo}Lean & Nati'


def anchos(lista):
    """Cuánto ocupa cada columna para que todo quede alineado.

    Se mide sobre los datos, no a ojo, para que al agregar un nombre largo
    la lista entera se reacomode sola en vez de quedar desprolija.
    """
    return (max(len(c) for c, _, _ in lista) + 5,
            max(len(n) for _, n, _ in lista) + 5)


def escribir_invitados_js(lista):
    cabecera = (RAIZ / 'invitados.js').read_text(encoding='utf-8').split('window.INVITADOS')[0]
    ancho_codigo, ancho_nombre = anchos(lista)
    filas = [
        "  {:<{a}} {{ nombre: {:<{b}} pases: {} }},".format(
            f"'{c}':", f"'{n}',", p, a=ancho_codigo, b=ancho_nombre)
        for c, n, p in lista
    ]
    (RAIZ / 'invitados.js').write_text(
        cabecera + 'window.INVITADOS = {\n' + '\n'.join(filas) + '\n};\n', encoding='utf-8')


def escribir_envios(lista):
    ancho_codigo, ancho_nombre = anchos(lista)
    filas = [
        "    [{:<{a}} {:<{b}} {}],".format(f"'{c}',", f"'{n}',", p,
                                           a=ancho_codigo, b=ancho_nombre)
        for c, n, p in lista
    ]
    html = (AQUI / 'envios.plantilla.html').read_text(encoding='utf-8')
    html = (html.replace('@@INVITADOS@@', '\n'.join(filas))
                .replace('@@TOTAL@@', str(len(lista)))
                .replace('@@LUGARES@@', str(sum(p for _, _, p in lista))))
    (SALIDA / 'envios.html').write_text(html, encoding='utf-8')


def escribir_planilla(lista):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Invitados'

    titulos = ['Invitación', 'Lugares', 'Código', 'Link personalizado', '¿Mandada?', '¿Confirmó?']
    hoja.append(titulos)

    relleno = PatternFill('solid', fgColor=BORDO)
    borde = Border(bottom=Side('thin', color=LINEA))
    for celda in hoja[1]:
        celda.font = Font(bold=True, color=PAPEL, size=11)
        celda.fill = relleno
        celda.alignment = Alignment(vertical='center')
    hoja.row_dimensions[1].height = 24

    for codigo, nombre, pases in lista:
        hoja.append([nombre, pases, codigo, f'{SITIO}?i={codigo}', '', ''])

    ultima = hoja.max_row
    for fila in hoja.iter_rows(min_row=2, max_row=ultima):
        for celda in fila:
            celda.border = borde
            celda.font = Font(color=TINTA, size=11)
        fila[1].alignment = Alignment(horizontal='center')
        fila[3].font = Font(color='1155CC', size=10, underline='single')
        fila[3].hyperlink = fila[3].value

    hoja.append([])
    hoja.append(['Totales', f'=SUM(B2:B{ultima})'])
    for celda in hoja[hoja.max_row][:2]:
        celda.font = Font(bold=True, color=BORDO, size=11)

    for columna, ancho in zip('ABCDEF', (30, 9, 26, 62, 12, 12)):
        hoja.column_dimensions[columna].width = ancho
    hoja.freeze_panes = 'A2'
    hoja.auto_filter.ref = f'A1:{get_column_letter(len(titulos))}{ultima}'

    libro.save(SALIDA / 'invitados-lean-y-nati.xlsx')


def main():
    lista = leer_lista()
    SALIDA.mkdir(exist_ok=True)
    escribir_invitados_js(lista)
    escribir_envios(lista)
    escribir_planilla(lista)
    print(f'{len(lista)} invitaciones · {sum(p for _, _, p in lista)} lugares')
    print('escritos: invitados.js · salida/envios.html · salida/invitados-lean-y-nati.xlsx')


if __name__ == '__main__':
    main()
